"""سخت‌سازی امنیتی: رمز و کد در لاگ نمی‌آید، فرمول در خروجی اکسل/CSV خنثی است، تنظیمات ناامن سرور اصلی رد می‌شود."""
import io
import logging
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.core import security
from app.core.http_security import RedactTokenPaths
from app.services import messaging
from app.services.spreadsheet_safety import csv_cell, neutralize_workbook

TEMP_PASSWORD = "Tmp-Secret-9921"
OTP = "482913"


@pytest.fixture
def channels_off(monkeypatch):
    monkeypatch.setattr(messaging, "sms_enabled", lambda: False)
    monkeypatch.setattr(messaging, "email_enabled", lambda: False)


def test_unconfigured_messaging_never_logs_message_body(channels_off, caplog):
    caplog.set_level(logging.DEBUG)
    request = SimpleNamespace(contact_mobile="09121231087", contact_email="owner@example.com", name="فروشگاه آزمون",
                              contact_name="علی", tracking_code="RQ-1")
    messaging.organization_request_approved(request, "shop_admin", TEMP_PASSWORD)
    user = SimpleNamespace(mobile="09121231087", email="owner@example.com")
    messaging.password_reset_code(user, OTP, "sms")
    messaging.password_reset_code(user, OTP, "email")
    messaging.deliver_sms("09121231087", f"رمز موقت: {TEMP_PASSWORD}")
    messaging.deliver_email("owner@example.com", "موضوع", f"کد: {OTP}")

    logged = caplog.text
    assert TEMP_PASSWORD not in logged and OTP not in logged
    assert "09121231087" not in logged and "owner@example.com" not in logged
    assert "0912***1087" in logged and "o***@example.com" in logged


@pytest.mark.parametrize("value,expected", [
    ("=HYPERLINK(\"http://x\")", "'=HYPERLINK(\"http://x\")"), ("+98 912", "'+98 912"), ("@SUM(A1)", "'@SUM(A1)"),
    ("\tcmd", "'\tcmd"), ("-2+3", "'-2+3"), ("-12.5", "-12.5"), ("-", "-"), ("کالا", "کالا"), (42, 42), (None, None),
])
def test_csv_cells_are_neutralized(value, expected):
    assert csv_cell(value) == expected


def test_workbook_formulas_are_stored_as_text():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "=1+1"
    ws["A2"] = "متن عادی"
    ws["A3"] = -5
    neutralize_workbook(wb)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    loaded = load_workbook(out).active
    assert loaded["A1"].data_type == "s" and loaded["A1"].value == "=1+1"
    assert loaded["A2"].value == "متن عادی" and loaded["A3"].value == -5


def test_production_rejects_placeholder_and_short_secrets(monkeypatch):
    monkeypatch.setenv("APP_ENV", "production")
    monkeypatch.setenv("DATABASE_URL", "postgresql://postgres:your_secure_password@db:5432/logitrack")
    for bad in ("", "change-me-to-a-long-random-string", "short-key"):
        monkeypatch.setenv("SECRET_KEY", bad)
        with pytest.raises(RuntimeError) as e:
            security.enforce_production_settings()
        assert bad not in str(e.value) or not bad      # مقدار کلید هرگز در پیام خطا نمی‌آید

    monkeypatch.setenv("SECRET_KEY", "k" * 48)
    monkeypatch.setenv("DATABASE_URL", "postgresql://postgres:a-real-random-db-password@db:5432/logitrack")
    monkeypatch.delenv("DEMO_DATABASE_URL", raising=False)
    security.enforce_production_settings()
    assert all(ok for _, ok in security.insecure_settings())


def test_development_keeps_working_without_secrets(monkeypatch):
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.setenv("SECRET_KEY", "")
    security.enforce_production_settings()


def test_access_log_hides_secret_link_tokens():
    record = logging.LogRecord("uvicorn.access", logging.INFO, __file__, 1, '%s - "%s %s HTTP/%s" %d',
                               ("1.2.3.4:5", "GET", "/app/d/demoTokenAbcdefghijklmnopqrstuvwxyz", "1.1", 200), None)
    RedactTokenPaths().filter(record)
    assert record.args[2] == "/app/d/***"
    for path in ("/support/abc123", "/public/tickets/abc123/messages", "/join/QAZVIN26", "/driver-signup/QAZ", "/rate/R1"):
        record.args = ("ip", "GET", path, "1.1", 200)
        RedactTokenPaths().filter(record)
        assert "abc123" not in record.args[2] and "QAZ" not in record.args[2] and "R1" not in record.args[2]
