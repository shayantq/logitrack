"""
مرحله ۱۳: کاتالوگ کالا — قیمت تاریخ‌دار با یارانه (سفارش قیمت روز تحویل را می‌گیرد)، تنوع، و ورود/خروج اکسل.
"""
import io
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import Workbook, load_workbook

from app.core.iran import gregorian_to_jalali
from app.core.security import create_access_token
from app.core.tenancy import set_tenant
from app.models import Inventory, OrderItem, Organization, Product, ProductPrice, RoleEnum, StockMovement, User, Vehicle, Warehouse
from app.services import catalog_io
from app.services.catalog_service import today_local
from tests.factories import order_payload


def _auth(user) -> dict:
    return {"Authorization": f"Bearer {create_access_token(user.username, user.id, timedelta(minutes=5))}"}


@pytest.fixture()
def env(api_env):
    client, Session = api_env
    db = Session()
    world = {}
    for slug in ("alpha", "beta"):
        org = Organization(name=f"سازمان {slug}", slug=slug)
        db.add(org)
        db.commit()
        set_tenant(db, org.id)
        tehran = Warehouse(name="انبار تهران", province="تهران", city="تهران")
        karaj = Warehouse(name="انبار کرج", province="البرز", city="کرج")
        db.add_all([tehran, karaj])
        db.commit()
        admin = User(username=f"{slug}_admin", email=f"{slug}a@x.io", role=RoleEnum.ORG_ADMIN, hashed_password="x")
        manager = User(username=f"{slug}_mgr", email=f"{slug}m@x.io", role=RoleEnum.WAREHOUSE_MANAGER, hashed_password="x", warehouse_id=tehran.id)
        driver = User(username=f"{slug}_d", email=f"{slug}d@x.io", role=RoleEnum.DRIVER, hashed_password="x", city="تهران")
        db.add_all([admin, manager, driver])
        db.commit()
        vehicle = Vehicle(license_plate=f"{slug}-1", type="van", capacity_kg=50_000, status="available", driver_id=driver.id, city="تهران")
        db.add(vehicle)
        db.commit()
        world[slug] = dict(org=org, tehran=tehran, karaj=karaj, admin=admin, manager=manager, driver=driver, vehicle=vehicle)
        set_tenant(db, None)
    try:
        yield client, db, world
    finally:
        db.close()


def _product(client, w, sku="TV-55", price=1_000_000, **extra):
    res = client.post("/inventory/products", json={"name": "تلویزیون ۵۵ اینچ", "sku": sku, "price": price, "weight_kg": 12,
                                                   "initial_quantity": 50, "warehouse_id": w["tehran"].id, **extra}, headers=_auth(w["admin"]))
    assert res.status_code == 201, res.text
    return res.json()["product"]["id"]


def _order(client, w, product_id, day_offset=0):
    start = datetime.now(timezone.utc) + timedelta(days=day_offset)
    return client.post("/orders/", json=order_payload(product_id, w["vehicle"].id, delivery_window_start=start.isoformat(),
                                                      delivery_window_end=(start + timedelta(hours=4)).isoformat()), headers=_auth(w["admin"]))


def _price(client, w, pid, days, base, subsidy=0, **extra):
    return client.post(f"/catalog/products/{pid}/prices", json={"effective_from": (today_local() + timedelta(days=days)).isoformat(),
                                                                 "base_price": base, "subsidy_percent": subsidy, **extra}, headers=_auth(w["admin"]))


# ---------------------------------------------------------------------------
def test_order_takes_price_of_delivery_day_and_past_orders_never_change(env):
    client, db, w = env
    a = w["alpha"]
    pid = _product(client, a)
    today_order = _order(client, a, pid)
    assert today_order.status_code == 201, today_order.text
    assert today_order.json()["items"][0]["unit_price"] == 1_000_000
    # قیمت جدید از ۵ روز بعد (مثلاً گران شدن دلار)
    assert _price(client, a, pid, 5, 1_300_000).status_code == 200
    later = _order(client, a, pid, day_offset=6).json()
    soon = _order(client, a, pid, day_offset=2).json()
    assert later["items"][0]["unit_price"] == 1_300_000 and later["total_price"] == 1_300_000
    assert soon["items"][0]["unit_price"] == 1_000_000
    detail = client.get(f"/catalog/products/{pid}", headers=_auth(a["admin"])).json()
    assert detail["current_price"] == 1_000_000 and [p["is_future"] for p in detail["prices"]] == [True, False]
    # حتی قیمت جدید از امروز، سفارش ثبت‌شده را عوض نمی‌کند
    _price(client, a, pid, 0, 900_000)
    set_tenant(db, a["org"].id)
    assert db.query(OrderItem).filter(OrderItem.order_id == today_order.json()["id"]).one().unit_price == 1_000_000
    set_tenant(db, None)


def test_subsidy_price_and_price_deletion_rules(env):
    client, db, w = env
    a = w["alpha"]
    pid = _product(client, a, sku="OIL", price=6_000_000, subsidy_percent=20)
    detail = client.get(f"/catalog/products/{pid}", headers=_auth(a["admin"])).json()
    assert (detail["current_base_price"], detail["current_subsidy_percent"], detail["current_price"]) == (6_000_000, 20, 4_800_000)
    listed = next(p for p in client.get("/inventory/products", headers=_auth(a["admin"])).json() if p["id"] == pid)
    assert listed["price"] == 4_800_000 and listed["subsidy_percent"] == 20
    assert _price(client, a, pid, 0, 1, subsidy=100).status_code == 422

    only = detail["prices"][0]["id"]
    assert client.delete(f"/catalog/prices/{only}", headers=_auth(a["admin"])).status_code == 400      # تنها قیمت
    future = _price(client, a, pid, 30, 7_000_000, 20).json()
    future_id = next(p["id"] for p in future["prices"] if p["is_future"])
    assert _order(client, a, pid).status_code == 201                                                     # قیمت فعلی استفاده شد
    assert client.delete(f"/catalog/prices/{future_id}", headers=_auth(a["admin"])).status_code == 200  # آینده: آزاد
    _price(client, a, pid, 10, 6_500_000)
    after = client.get(f"/catalog/products/{pid}", headers=_auth(a["admin"])).json()
    current_id = next(p["id"] for p in after["prices"] if p["is_current"])
    assert client.delete(f"/catalog/prices/{current_id}", headers=_auth(a["admin"])).status_code == 400  # در سفارش استفاده شده
    assert client.post(f"/catalog/products/{pid}/prices", json={"effective_from": today_local().isoformat(), "base_price": 1},
                       headers=_auth(a["driver"])).status_code == 403


def test_variants_and_detailed_fields(env):
    client, db, w = env
    a = w["alpha"]
    pid = _product(client, a, sku="PS5-AE", price=480_000_000, origin="imported", origin_country="امارات", category="کنسول بازی")
    upd = client.put(f"/catalog/products/{pid}", json={"specs": [["رنگ", "سفید"], ["گارانتی", "۱۸ ماه"]], "hs_code": "95045000",
                                                       "export_eligible": True, "brand": "Sony"}, headers=_auth(a["manager"]))
    assert upd.status_code == 200 and upd.json()["specs"][1] == ["گارانتی", "۱۸ ماه"]
    var = client.post(f"/catalog/products/{pid}/variants", json={"sku": "PS5-CN", "variant_label": "سری واردات چین", "origin": "imported",
                                                                 "origin_country": "چین", "base_price": 455_000_000}, headers=_auth(a["admin"]))
    assert var.status_code == 201, var.text
    v = var.json()
    assert v["parent"]["sku"] == "PS5-AE" and v["origin_country"] == "چین" and v["current_price"] == 455_000_000
    assert v["specs"][0] == ["رنگ", "سفید"] and v["brand"] == "Sony" and v["stock"][0]["warehouse_name"] == "انبار تهران"
    parent = client.get(f"/catalog/products/{pid}", headers=_auth(a["admin"])).json()
    assert [x["sku"] for x in parent["variants"]] == ["PS5-CN"]
    # تنوعِ تنوع ← به کالای اصلی وصل می‌شود؛ کد تکراری رد می‌شود
    v2 = client.post(f"/catalog/products/{v['id']}/variants", json={"sku": "PS5-JP", "variant_label": "ژاپن", "base_price": 1},
                     headers=_auth(a["admin"])).json()
    assert v2["parent"]["sku"] == "PS5-AE"
    assert client.post(f"/catalog/products/{pid}/variants", json={"sku": "PS5-JP", "variant_label": "سری دیگر", "base_price": 1},
                       headers=_auth(a["admin"])).status_code == 400
    # کالای دیگر سازمان دیده نمی‌شود
    assert client.get(f"/catalog/products/{pid}", headers=_auth(w["beta"]["admin"])).status_code == 404


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "کالاها"
    ws.append([title + (" *" if req else "") for _, title, req, *_ in catalog_io.COLUMNS])
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _row(**values):
    row = [None] * len(catalog_io.KEYS)
    for k, v in values.items():
        row[catalog_io.KEYS.index(k)] = v
    return row


def _import(client, user, content, stock_mode="ignore", dry_run=True):
    return client.post("/catalog/import", data={"stock_mode": stock_mode, "dry_run": str(dry_run).lower()},
                       files={"file": ("c.xlsx", content, catalog_io.__name__)}, headers=_auth(user))


def test_excel_template_export_and_import(env):
    client, db, w = env
    a = w["alpha"]
    tpl = client.get("/catalog/template", headers=_auth(a["admin"]))
    assert tpl.status_code == 200 and "spreadsheetml" in tpl.headers["content-type"]
    book = load_workbook(io.BytesIO(tpl.content))
    assert book.sheetnames == ["کالاها", "راهنما", "مقادیر"] and book["کالاها"].max_row == 4

    jy, jm, jd = gregorian_to_jalali(*(today_local() + timedelta(days=10)).timetuple()[:3])
    bad = _xlsx([
        _row(sku="X1", name="کالا", price="abc"),
        _row(sku="X2", name="کالا ۲", price=1000, origin="مریخی", warehouse="انبار ناموجود", quantity=5),
        _row(sku="X3", name="کالا ۳"),
    ])
    report = _import(client, a["admin"], bad, stock_mode="set").json()
    assert report["error_count"] == 3 and report["applied"] is False
    assert _import(client, a["admin"], bad, stock_mode="set", dry_run=False).json()["applied"] is False   # همه یا هیچ

    good = _xlsx([
        _row(sku="RICE", name="برنج ۱۰ کیلویی", price="۱٬۸۵۰٬۰۰۰", subsidy_percent=10, unit="کیسه", category="مواد غذایی",
             origin="وارداتی", origin_country="هند", specs="دانه: بلند؛ وزن: ۱۰ کیلو", warehouse="انبار تهران", quantity=40, reorder_point=5),
        _row(sku="RICE", warehouse="انبار کرج", quantity=15),
        _row(sku="RICE-PK", name="برنج ۱۰ کیلویی", price=1_700_000, parent_sku="RICE", variant_label="سری پاکستان",
             origin="وارداتی", origin_country="پاکستان", effective_from=f"{jy}/{jm:02d}/{jd:02d}", warehouse="انبار تهران", quantity=10),
    ])
    preview = _import(client, a["admin"], good, stock_mode="set").json()
    assert (preview["error_count"], preview["create"], preview["stock_rows"]) == (0, 2, 3), preview
    done = _import(client, a["admin"], good, stock_mode="set", dry_run=False).json()
    assert done["applied"] is True
    set_tenant(db, a["org"].id)
    rice = db.query(Product).filter(Product.sku == "RICE").one()
    pk = db.query(Product).filter(Product.sku == "RICE-PK").one()
    assert rice.price == 1_665_000 and rice.unit == "کیسه" and rice.specs == [["دانه", "بلند"], ["وزن", "۱۰ کیلو"]]
    assert pk.parent_id == rice.id and pk.price == 0                         # قیمتش از ۱۰ روز بعد اعمال می‌شود
    assert db.query(ProductPrice).filter(ProductPrice.product_id == pk.id).one().effective_from == today_local() + timedelta(days=10)
    stock = {(i.product_id, i.warehouse_id): i.quantity for i in db.query(Inventory).all()}
    assert stock[(rice.id, a["tehran"].id)] == 40 and stock[(rice.id, a["karaj"].id)] == 15
    assert db.query(StockMovement).filter(StockMovement.reason == "تنظیم موجودی از فایل اکسل").count() == 3
    set_tenant(db, None)

    # خروجی ← ورود دوباره بدون خطا (فقط به‌روزرسانی)؛ «افزودن به موجودی» موجودی را زیاد می‌کند
    exported = client.get("/catalog/export", headers=_auth(a["admin"])).content
    again = _import(client, a["admin"], exported).json()
    assert again["error_count"] == 0 and again["create"] == 0 and again["update"] == 2
    add = _xlsx([_row(sku="RICE", warehouse="انبار تهران", quantity=5)])
    assert _import(client, a["admin"], add, stock_mode="add", dry_run=False).json()["applied"] is True
    set_tenant(db, a["org"].id)
    db.expire_all()
    assert db.query(Inventory).filter(Inventory.product_id == rice.id, Inventory.warehouse_id == a["tehran"].id).one().quantity == 45
    set_tenant(db, None)
    # مدیر انبار فقط انبار خودش
    other = _xlsx([_row(sku="RICE", warehouse="انبار کرج", quantity=1)])
    assert _import(client, a["manager"], other, stock_mode="add").json()["error_count"] == 1
    assert _import(client, a["driver"], other).status_code == 403
