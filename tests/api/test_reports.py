"""
مرحله ۱۴: گزارش‌های اکسل/PDF با برند سازمان و داشبورد مالی حمل‌ونقل.
"""
import base64
import io
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.core import storage
from app.core.security import create_access_token
from app.core.tenancy import set_tenant
from app.models import Order, Organization, RoleEnum, User, Vehicle, Warehouse
from app.routers import reports as reports_router
from app.services import report_data
from app.services.catalog_service import today_local
from app.tasks import build_report_file
from tests.factories import order_payload


def _auth(user) -> dict:
    return {"Authorization": f"Bearer {create_access_token(user.username, user.id, timedelta(minutes=5))}"}


@pytest.fixture()
def env(api_env, monkeypatch, tmp_path):
    monkeypatch.setattr(storage, "MEDIA_ROOT", tmp_path)
    client, Session = api_env
    db = Session()
    world = {}
    for slug in ("alpha", "beta"):
        org = Organization(name=f"سازمان {slug}", slug=slug)
        db.add(org)
        db.commit()
        set_tenant(db, org.id)
        tehran = Warehouse(name="انبار تهران", province="تهران", city="تهران")
        db.add(tehran)
        db.commit()
        admin = User(username=f"{slug}_admin", email=f"{slug}a@x.io", role=RoleEnum.ORG_ADMIN, hashed_password="x")
        manager = User(username=f"{slug}_mgr", email=f"{slug}m@x.io", role=RoleEnum.WAREHOUSE_MANAGER, hashed_password="x", warehouse_id=tehran.id)
        driver = User(username=f"{slug}_d", email=f"{slug}d@x.io", role=RoleEnum.DRIVER, hashed_password="x", city="تهران",
                      full_name=f"راننده {slug}")
        db.add_all([admin, manager, driver])
        db.commit()
        vehicle = Vehicle(license_plate=f"{slug}-1", type="van", capacity_kg=50_000, status="available", driver_id=driver.id, city="تهران")
        db.add(vehicle)
        db.commit()
        world[slug] = dict(org=org, tehran=tehran, admin=admin, manager=manager, driver=driver, vehicle=vehicle)
        set_tenant(db, None)
    try:
        yield client, db, world
    finally:
        db.close()


def _delivered_orders(client, db, w, fares=(1_000_000, 3_000_000), city="تهران"):
    res = client.post("/inventory/products", json={"name": "برنج ۱۰ کیلویی", "sku": f"R-{w['org'].slug}", "price": 500_000, "weight_kg": 10,
                                                   "initial_quantity": 100, "warehouse_id": w["tehran"].id}, headers=_auth(w["admin"]))
    assert res.status_code == 201, res.text
    pid = res.json()["product"]["id"]
    now = datetime.now(timezone.utc)
    ids = []
    for _ in fares:
        r = client.post("/orders/", json=order_payload(pid, w["vehicle"].id, city=city, delivery_window_start=now.isoformat(),
                                                       delivery_window_end=(now + timedelta(hours=4)).isoformat()), headers=_auth(w["admin"]))
        assert r.status_code == 201, r.text
        ids.append(r.json()["id"])
    set_tenant(db, None)
    for i, (oid, fare) in enumerate(zip(ids, fares, strict=False)):
        o = db.get(Order, oid)
        o.status = "delivered"
        o.driver_id = w["driver"].id
        o.delivered_at = now.replace(tzinfo=None) - timedelta(minutes=5)
        o.delivery_fee = fare
        o.commission_percent = 20
        o.commission_amount = fare * 0.2
        o.driver_earning = fare * 0.8
        o.commission_level = "L1"
        o.day_trip_number = i + 1
        o.route = {"total_km": 100}
    db.commit()
    return ids


# ---------------------------------------------------------------------------
def test_finance_dashboard_math_and_tenant_isolation(env):
    client, db, world = env
    a, b = world["alpha"], world["beta"]
    _delivered_orders(client, db, a)
    _delivered_orders(client, db, b, fares=(9_000_000,))
    res = client.get("/reports/finance", headers=_auth(a["admin"]))
    assert res.status_code == 200, res.text
    data = res.json()
    k = data["kpis"]
    assert (k["trips"], k["fare"], k["company"], k["driver"], k["km"]) == (2, 4_000_000, 800_000, 3_200_000, 200)
    assert (k["avg_fare"], k["fare_per_km"], k["take_rate"]) == (2_000_000, 20_000, 20.0)
    assert [d["driver"] for d in data["by_driver"]] == ["راننده alpha"]
    assert data["by_driver"][0]["driver_share"] == 3_200_000
    assert data["by_province"][0]["province"] == "تهران" and data["by_province"][0]["trips"] == 2
    assert sum(m["fare"] for m in data["by_month"]) == 4_000_000
    assert len(data["details"]) == 2 and all(d["fare"] < 9_000_000 for d in data["details"])
    # بازه خالی در گذشته: صفر
    old = today_local() - timedelta(days=400)
    empty = client.get("/reports/finance", params={"date_from": old.isoformat(), "date_to": (old + timedelta(days=10)).isoformat()},
                       headers=_auth(a["admin"])).json()
    assert empty["kpis"]["trips"] == 0 and empty["kpis"]["take_rate"] == 0
    bad = client.get("/reports/finance", params={"date_from": today_local().isoformat(),
                                                 "date_to": (today_local() - timedelta(days=1)).isoformat()}, headers=_auth(a["admin"]))
    assert bad.status_code == 400


def test_report_access_rules(env, monkeypatch):
    client, db, world = env
    a = world["alpha"]
    queued = []

    class _Task:
        id = "task-1"

    monkeypatch.setattr(reports_router.generate_report_task, "delay", lambda *args: queued.append(args) or _Task())
    # مدیر انبار: مالی ممنوع، موجودی مجاز ولی محدود به انبار خودش
    assert client.get("/reports/finance", headers=_auth(a["manager"])).status_code == 403
    assert client.post("/reports/generate", json={"report": "finance"}, headers=_auth(a["manager"])).status_code == 403
    assert client.post("/reports/generate", json={"report": "finance"}, headers=_auth(a["driver"])).status_code == 403
    ok = client.post("/reports/generate", json={"report": "ledger", "format": "pdf", "date_from": "2026-01-01"}, headers=_auth(a["manager"]))
    assert ok.status_code == 202, ok.text
    assert ok.json()["task_id"] == "task-1"
    org_id, user_id, report, fmt, params = queued[-1]
    assert (org_id, user_id, report, fmt, params["date_from"]) == (a["org"].id, a["manager"].id, "ledger", "pdf", "2026-01-01")
    bad = client.post("/reports/generate", json={"report": "orders", "date_from": "2026-02-01", "date_to": "2026-01-01"}, headers=_auth(a["admin"]))
    assert bad.status_code == 400
    assert client.post("/reports/generate", json={"report": "salary"}, headers=_auth(a["admin"])).status_code == 422


def test_branded_files_are_valid_and_carry_brand(env):
    client, db, world = env
    a = world["alpha"]
    _delivered_orders(client, db, a)
    # لوگو و رنگ سازمان
    png = io.BytesIO()
    Image.new("RGB", (120, 60), (200, 30, 30)).save(png, "PNG")
    assert client.post("/organization/logo", files={"file": ("logo.png", png.getvalue(), "image/png")}, headers=_auth(a["admin"])).status_code == 200
    prof = client.put("/organization/profile", json={"brand_color": "#B91C1C", "legal_name": "شرکت آلفا"}, headers=_auth(a["admin"]))
    assert prof.status_code == 200, prof.text
    assert prof.json()["brand_color"] == "#B91C1C"
    assert client.put("/organization/profile", json={"brand_color": "red"}, headers=_auth(a["admin"])).status_code == 422

    set_tenant(db, a["org"].id)
    admin = db.get(User, a["admin"].id)
    for report in report_data.REPORTS:
        for fmt in ("excel", "pdf"):
            out = build_report_file(db, a["org"].id, admin, report, fmt, {"date_from": (today_local() - timedelta(days=3)).isoformat()})
            content = base64.b64decode(out["content_b64"])
            assert out["organization_id"] == a["org"].id
            if fmt == "pdf":
                assert content.startswith(b"%PDF") and out["filename"].endswith(".pdf")
            else:
                assert out["filename"].endswith(".xlsx")
                wb = load_workbook(io.BytesIO(content))
                ws = wb.worksheets[0]
                assert ws.sheet_view.rightToLeft
                texts = [str(c.value) for row in ws.iter_rows(max_row=12) for c in row if c.value is not None]
                assert any("سازمان alpha" in t for t in texts)
                assert ws._images, "لوگو در اکسل نیست"
                if report == "finance":
                    assert len(wb.worksheets) == 4
    # خروجی سریع صفحه سفارش‌ها (دکمه‌های قدیمی) در سقف بازه سه‌ساله می‌ماند و سفارش امروز را دارد
    quick = build_report_file(db, a["org"].id, admin, "orders", "excel", reports_router._all_orders_params())
    rows = [r for r in load_workbook(io.BytesIO(base64.b64decode(quick["content_b64"]))).worksheets[0].iter_rows(values_only=True)]
    assert sum(1 for r in rows if any(isinstance(c, str) and c.startswith("LT-") for c in r)) == 2
    # مدیر انبار فقط انبار خودش را در گزارش موجودی می‌بیند
    other = Warehouse(name="انبار دوم", province="البرز", city="کرج")
    db.add(other)
    db.commit()
    manager = db.get(User, a["manager"].id)
    rep = report_data.inventory(db, manager, warehouse_id=other.id)
    assert rep.sections[0].rows and all(r["warehouse"] == "انبار تهران" for r in rep.sections[0].rows)
    set_tenant(db, None)
