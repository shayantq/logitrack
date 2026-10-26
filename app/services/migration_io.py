"""
ابزار مهاجرت داده برای راه‌اندازی (فصل‌های ۳۲ و ۳۳): ورود «محل‌های انبار» و «موجودی اول دوره» از اکسل.

هر دو مثل ورود کاتالوگ دو مرحله دارند: پیش‌نمایش (اعتبارسنجی همه ردیف‌ها بدون ذخیره) و ثبت نهایی (همه یا هیچ).
موجودی اول دوره یک سند «موجودی اول دوره» (OB) با ردیف‌های ریز محل/بچ/سریال/انقضا می‌سازد؛ ظرفیت و قاعده اختلاط
محل‌ها، ردیابی بچ/سریال/انقضای کالا و یکتایی سریال همان قاعده‌های دفتر گردش‌اند. فایل تکراری (همان محتوا) دوباره
ثبت نمی‌شود (کلید یکتایی سند = اثر انگشت فایل).
"""
from __future__ import annotations

import hashlib
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.iran import normalize_persian_text
from app.core.wms import LOCATION_LEVEL_LABELS, LOCATION_LEVELS, LOCATION_TYPES, STOCK_STATUS_LABELS, STOCK_STATUSES
from app.models import Location, Product, StockDocument, User, Warehouse
from app.services import audit_service, location_service, permission_service, stock_ledger
from app.services.catalog_io import _date, _number, _text
from app.services.exceptions import ConflictError, DomainError

MAX_ROWS = 5000
KINDS = {"locations": "محل‌های انبار", "opening": "موجودی اول دوره"}
COLUMNS = {
    "locations": [("warehouse", "انبار", True, 22), ("code", "کد محل", True, 16), ("level", "سطح", True, 14),
                  ("parent", "کد محل مادر", False, 16), ("type", "نوع محل", False, 16), ("barcode", "بارکد", False, 18),
                  ("max_weight_kg", "حداکثر وزن (کیلوگرم)", False, 14), ("max_volume_m3", "حداکثر حجم (مترمکعب)", False, 14),
                  ("max_pallets", "حداکثر پالت", False, 12), ("route_sequence", "ترتیب مسیر", False, 12)],
    "opening": [("warehouse", "انبار", True, 22), ("sku", "کد کالا", True, 16), ("quantity", "مقدار", True, 10),
                ("location", "کد محل", False, 14), ("lot", "بچ", False, 14), ("expiry", "تاریخ انقضا", False, 14),
                ("serial", "سریال", False, 16), ("status", "وضعیت کیفی", False, 16)],
}
SAMPLES = {
    "locations": [["{wh}", "A", "منطقه", "", "نگهداری", "", "", "", "", ""],
                  ["{wh}", "A-01", "راهرو", "A", "نگهداری", "", "", "", "", ""],
                  ["{wh}", "A-01-01", "خانه / جایگاه", "A-01", "برداشت", "", 500, 2, 2, 1],
                  ["{wh}", "A-01-02", "خانه / جایگاه", "A-01", "نگهداری", "", 1000, 4, 4, 2]],
    "opening": [["{wh}", "{sku}", 120, "A-01-01", "", "", "", "سالم و آزاد"]],
}
LEVELS = {**{v: k for k, v in LOCATION_LEVEL_LABELS.items()}, **{k: k for k in LOCATION_LEVELS}, "خانه": "bin"}
TYPES = {**{t.label: t.code for t in LOCATION_TYPES.values()}, **{c: c for c in LOCATION_TYPES}}
STATUSES = {**{v: k for k, v in STOCK_STATUS_LABELS.items()}, **{s: s for s in STOCK_STATUSES}, "سالم": "available"}


def _ensure(db: Session, user: User) -> None:
    location_service.ensure_staff(user)
    permission_service.ensure(db, user, "data.import")


def template(db: Session, user: User, kind: str) -> bytes:
    _ensure(db, user)
    if kind not in KINDS:
        raise ConflictError("نوع قالب معتبر نیست.")
    warehouses = [w.name for w in db.query(Warehouse).order_by(Warehouse.id).all()]
    product = db.query(Product).order_by(Product.id).first()
    wb = Workbook()
    ws = wb.active
    ws.title = KINDS[kind]
    ws.sheet_view.rightToLeft = True
    head = PatternFill("solid", fgColor="0D9488")
    for i, (_, title, required, width) in enumerate(COLUMNS[kind], start=1):
        cell = ws.cell(row=1, column=i, value=title + (" *" if required else ""))
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    values = wb.create_sheet("مقادیر")
    values.sheet_view.rightToLeft = True
    lists = [("انبار", warehouses), ("سطح", [LOCATION_LEVEL_LABELS[lv] for lv in LOCATION_LEVELS]),
             ("نوع محل", [t.label for t in LOCATION_TYPES.values()]), ("وضعیت کیفی", [STOCK_STATUS_LABELS[s] for s in STOCK_STATUSES])]
    for col, (title, items) in enumerate(lists, start=1):
        values.cell(row=1, column=col, value=title).font = Font(bold=True)
        for r, v in enumerate(items, start=2):
            values.cell(row=r, column=col, value=v)
        values.column_dimensions[values.cell(row=1, column=col).column_letter].width = 24
    keys = [c[0] for c in COLUMNS[kind]]
    for key, col in (("warehouse", "A"), ("level", "B"), ("type", "C"), ("status", "D")):
        if key in keys and lists["ABCD".index(col)][1]:
            letter = ws.cell(row=1, column=keys.index(key) + 1).column_letter
            dv = DataValidation(type="list", formula1=f"=مقادیر!${col}$2:${col}${len(lists['ABCD'.index(col)][1]) + 1}", allow_blank=True)
            dv.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
            ws.add_data_validation(dv)
    for r, row in enumerate(SAMPLES[kind], start=2):
        for c, v in enumerate(row, start=1):
            if isinstance(v, str):
                v = v.replace("{wh}", warehouses[0] if warehouses else "").replace("{sku}", product.sku if product else "")
            ws.cell(row=r, column=c, value=v)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _read(content: bytes, kind: str) -> list[tuple[int, dict]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ConflictError("فایل اکسل خوانده نشد؛ قالب ‎.xlsx‎ سامانه را دانلود و پر کنید.")
    ws = wb[KINDS[kind]] if KINDS[kind] in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    titles = {re.sub(r"\s*\*$", "", str(h or "")).strip(): i for i, h in enumerate(header)}
    missing = [t for _, t, req, _ in COLUMNS[kind] if req and t not in titles]
    if missing:
        raise ConflictError(f"ستون‌های اجباری در فایل نیست: {'، '.join(missing)}. از قالب سامانه استفاده کنید.")
    index = {k: titles.get(t) for k, t, *_ in COLUMNS[kind]}
    result = []
    for n, values in enumerate(rows, start=2):
        if not values or all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        if len(result) >= MAX_ROWS:
            raise ConflictError(f"حداکثر {MAX_ROWS} ردیف در هر بارگذاری.")
        result.append((n, {k: (values[i] if i is not None and i < len(values) else None) for k, i in index.items()}))
    if not result:
        raise ConflictError("فایل هیچ ردیفی ندارد.")
    return result


def _warehouse_map(db: Session, user: User) -> dict[str, Warehouse]:
    result = {}
    for w in db.query(Warehouse).all():
        if getattr(user.role, "value", user.role) == "warehouse_manager" and w.id != user.warehouse_id:
            continue
        result[normalize_persian_text(w.name).strip()] = w
        if w.code:
            result[w.code.strip().upper()] = w
    return result


def _report(kind: str, rows: list[dict], applied: bool, extra: dict | None = None) -> dict:
    errors = [{"line": r["line"], "message": m} for r in rows for m in r["errors"]]
    return {"kind": kind, "rows": len(rows), "valid": sum(1 for r in rows if not r["errors"]), "error_count": len(errors),
            "errors": errors[:200], "applied": applied, "preview": [{"line": r["line"], "summary": r["summary"], "ok": not r["errors"]}
                                                               for r in rows[:50]], **(extra or {})}


def import_locations(db: Session, user: User, content: bytes, *, dry_run: bool = True) -> dict:
    _ensure(db, user)
    permission_service.ensure(db, user, "masterdata.locations")
    warehouses = _warehouse_map(db, user)
    rows = []
    file_codes: dict[tuple[int, str], dict] = {}
    for line, d in _read(content, "locations"):
        errors: list[str] = []
        wh = warehouses.get(normalize_persian_text(_text(d["warehouse"]) or "").strip()) or warehouses.get((_text(d["warehouse"]) or "").upper())
        if wh is None:
            errors.append(f"انبار «{_text(d['warehouse'])}» یافت نشد (یا به آن دسترسی ندارید).")
        code = (_text(d["code"]) or "").upper()
        if not code or not re.match(r"^[0-9A-Z._/-]{1,32}$", code):
            errors.append("«کد محل» خالی یا نامعتبر است (فقط حروف لاتین، عدد و - . _ /).")
        level = LEVELS.get(_text(d["level"]) or "")
        if level is None:
            errors.append(f"سطح «{_text(d['level'])}» معتبر نیست.")
        loc_type = TYPES.get(_text(d["type"]) or "نگهداری")
        if loc_type is None:
            errors.append(f"نوع محل «{_text(d['type'])}» معتبر نیست.")
        row = {"line": line, "errors": errors, "warehouse": wh, "code": code, "level": level, "type": loc_type,
               "parent": (_text(d["parent"]) or "").upper() or None, "barcode": _text(d["barcode"]),
               "max_weight_kg": _number(d["max_weight_kg"], "حداکثر وزن", errors, low=0.001),
               "max_volume_m3": _number(d["max_volume_m3"], "حداکثر حجم", errors, low=0.0001),
               "max_pallets": _number(d["max_pallets"], "حداکثر پالت", errors, integer=True, low=1),
               "route_sequence": _number(d["route_sequence"], "ترتیب مسیر", errors, integer=True, low=0),
               "summary": f"{code} — {LOCATION_LEVEL_LABELS.get(level, '?')} در {wh.name if wh else '?'}"}
        if wh is not None and code:
            if (wh.id, code) in file_codes:
                errors.append(f"کد محل {code} در فایل تکراری است.")
            file_codes[(wh.id, code)] = row
            if db.query(Location.id).filter(Location.warehouse_id == wh.id, func.upper(Location.code) == code).first():
                row["exists"] = True
        rows.append(row)
    for row in rows:
        if row["parent"] and row["warehouse"] is not None and (row["warehouse"].id, row["parent"]) not in file_codes and not \
                db.query(Location.id).filter(Location.warehouse_id == row["warehouse"].id, func.upper(Location.code) == row["parent"]).first():
            row["errors"].append(f"محل مادر {row['parent']} نه در فایل هست نه در انبار.")
    created = sum(1 for r in rows if not r["errors"] and not r.get("exists"))
    extra = {"create": created, "skip": sum(1 for r in rows if r.get("exists"))}
    if dry_run or any(r["errors"] for r in rows):
        return _report("locations", rows, False, extra)
    rank = {lv: i for i, lv in enumerate(LOCATION_LEVELS)}
    try:
        for row in sorted((r for r in rows if not r.get("exists")), key=lambda r: rank[r["level"]]):
            parent_id = None
            if row["parent"]:
                parent = db.query(Location).filter(Location.warehouse_id == row["warehouse"].id,
                                                   func.upper(Location.code) == row["parent"]).first()
                parent_id = parent.id
            data = location_service._validate(db, row["warehouse"], {
                "code": row["code"], "level": row["level"], "parent_id": parent_id, "location_type": row["type"],
                "barcode": row["barcode"], "max_weight_kg": row["max_weight_kg"], "max_volume_m3": row["max_volume_m3"],
                "max_pallets": row["max_pallets"], "route_sequence": row["route_sequence"], "is_blocked": False})
            db.add(Location(warehouse_id=row["warehouse"].id, **{k: data.get(k) for k in location_service.FIELDS if k in data}))
            db.flush()
        audit_service.log(db, user, "data.import", f"ورود {created} محل انبار از اکسل", target_type="import",
                          details={"kind": "locations", "created": created, "skipped": extra["skip"]})
        db.commit()
    except DomainError as e:
        db.rollback()
        raise ConflictError(f"ثبت متوقف شد: {e.message}")
    except Exception:
        db.rollback()
        raise
    return _report("locations", rows, True, extra)


def import_opening(db: Session, user: User, content: bytes, *, dry_run: bool = True) -> dict:
    _ensure(db, user)
    fingerprint = "import-" + hashlib.sha256(content).hexdigest()[:40]
    previous = db.query(StockDocument).filter(StockDocument.idempotency_key == fingerprint).first()
    if previous is not None:
        raise ConflictError(f"همین فایل قبلاً با سند {previous.code} ثبت شده است.")
    warehouses = _warehouse_map(db, user)
    products = {p.sku.upper(): p for p in db.query(Product).all()}
    rows = []
    serials: set[tuple[int, str]] = set()
    for line, d in _read(content, "opening"):
        errors: list[str] = []
        wh = warehouses.get(normalize_persian_text(_text(d["warehouse"]) or "").strip()) or warehouses.get((_text(d["warehouse"]) or "").upper())
        if wh is None:
            errors.append(f"انبار «{_text(d['warehouse'])}» یافت نشد (یا به آن دسترسی ندارید).")
        product = products.get((_text(d["sku"]) or "").upper())
        if product is None:
            errors.append(f"کالا با کد «{_text(d['sku'])}» در کاتالوگ نیست؛ اول کاتالوگ را وارد کنید.")
        qty = _number(d["quantity"], "مقدار", errors, integer=True, low=1, high=10_000_000)
        if qty is None and not any("مقدار" in e for e in errors):
            errors.append("«مقدار» خالی است.")
        location = None
        code = (_text(d["location"]) or "").upper()
        if code and wh is not None:
            location = db.query(Location).filter(Location.warehouse_id == wh.id,
                                                 (func.upper(Location.code) == code) | (func.upper(Location.barcode) == code)).first()
            if location is None:
                errors.append(f"محل {code} در «{wh.name}» نیست.")
            elif location_service.has_children(db, location.id):
                errors.append(f"محل {code} زیرمجموعه دارد؛ کالا فقط در خانه انتهایی می‌نشیند.")
        lot, serial = _text(d["lot"]), _text(d["serial"])
        expiry = _date(d["expiry"], errors)
        status = STATUSES.get(_text(d["status"]) or "available")
        if status is None:
            errors.append(f"وضعیت «{_text(d['status'])}» معتبر نیست.")
        if product is not None:
            if product.lot_tracked and not lot:
                errors.append(f"کالای {product.sku} بچ‌دار است؛ «بچ» را بنویسید.")
            if product.expiry_tracked and not expiry:
                errors.append(f"کالای {product.sku} تاریخ انقضا دارد؛ «تاریخ انقضا» را بنویسید.")
            if product.serial_tracked:
                if not serial:
                    errors.append(f"کالای {product.sku} سریالی است؛ هر سریال یک ردیف با مقدار ۱.")
                elif qty not in (None, 1):
                    errors.append("ردیف سریالی باید مقدار ۱ داشته باشد.")
                elif (product.id, serial) in serials:
                    errors.append(f"سریال {serial} در فایل تکراری است.")
                serials.add((product.id, serial or ""))
        rows.append({"line": line, "errors": errors, "warehouse": wh, "product": product, "quantity": qty, "location": location,
                     "lot": lot, "serial": serial if product is not None and product.serial_tracked else None,
                     "expiry": expiry, "status": status,
                     "summary": f"{product.sku if product else '?'} × {qty or '?'} در {wh.name if wh else '?'}"
                                f"{' / ' + location.code if location else ''}"})
    total_units = sum(r["quantity"] or 0 for r in rows if not r["errors"])
    extra = {"units": total_units, "document_code": None}
    if dry_run or any(r["errors"] for r in rows):
        return _report("opening", rows, False, extra)
    try:
        documents: dict[int, StockDocument] = {}
        for row in rows:
            wh = row["warehouse"]
            if wh.id not in documents:
                documents[wh.id] = stock_ledger.new_document(
                    db, "opening", organization_id=user.organization_id, user=user, warehouse_id=wh.id,
                    note="موجودی اول دوره از فایل اکسل (راه‌اندازی)",
                    idempotency_key=fingerprint if not documents else None)
            inv = stock_ledger.lock_bucket(db, row["product"].id, wh.id, row["status"], organization_id=user.organization_id)
            key = stock_ledger.StockKey(location_id=row["location"].id if row["location"] else None, lot_number=row["lot"],
                                        serial_number=row["serial"], expiry_date=row["expiry"])
            stock_ledger.record(db, inv, row["quantity"], "initial", document=documents[wh.id], user=user,
                                reason="موجودی اول دوره (راه‌اندازی)", lines=[(key, row["quantity"])])
        codes = [d.code for d in documents.values()]
        extra["document_code"] = "، ".join(codes)
        audit_service.log(db, user, "data.import", f"موجودی اول دوره ({total_units} واحد) از اکسل ثبت شد: {extra['document_code']}",
                          target_type="import", details={"kind": "opening", "rows": len(rows), "units": total_units, "documents": codes})
        db.commit()
    except DomainError as e:
        db.rollback()
        raise ConflictError(f"ثبت متوقف شد: {e.message}")
    except Exception:
        db.rollback()
        raise
    return _report("opening", rows, True, extra)
