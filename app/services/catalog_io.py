"""
ورود و خروج کاتالوگ با اکسل (مرحله ۱۳).

قالب: یک برگه «کالاها» (هر ردیف یک کالا در یک انبار؛ برای چند انبار، ردیف تکراری با همان کد کالا)،
برگه «راهنما» و برگه «مقادیر» (فهرست انبارها، واحدها، مبدأ — برای فهرست کشویی سلول‌ها).
ورود دو مرحله دارد: پیش‌نمایش (بدون ذخیره، با خطای هر ردیف) و ثبت نهایی (همه یا هیچ).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from app.core.iran import gregorian_to_jalali, jalali_to_gregorian, normalize_persian_text
from app.models import Inventory, Product, User, Warehouse
from app.services import catalog_service, stock_ledger
from app.services.exceptions import ConflictError, ForbiddenError

MAX_ROWS = 5000
UNITS = ["عدد", "کیلوگرم", "گرم", "لیتر", "متر", "بسته", "کارتن", "جفت", "دستگاه"]
ORIGIN_LABELS = {"domestic": "داخلی", "imported": "وارداتی"}

# (کلید، عنوان ستون، اجباری، عرض، توضیح)
COLUMNS = [
    ("sku", "کد کالا", True, 16, "یکتا در سازمان؛ اگر از قبل باشد، کالا به‌روز می‌شود"),
    ("name", "نام کالا", True, 32, ""),
    ("category", "گروه کالا", False, 18, "مثلاً لوازم برقی"),
    ("brand", "برند", False, 14, ""),
    ("unit", "واحد", False, 10, "از فهرست: " + "، ".join(UNITS)),
    ("barcode", "بارکد", False, 16, ""),
    ("weight_kg", "وزن (کیلوگرم)", False, 12, "پیش‌فرض ۱"),
    ("price", "قیمت پایه (ریال)", True, 16, "قیمت فروش پیش از یارانه"),
    ("subsidy_percent", "درصد یارانه", False, 11, "۰ تا ۹۹؛ قیمت نهایی = قیمت پایه × (۱ − درصد/۱۰۰)"),
    ("cost_price", "بهای تمام‌شده (ریال)", False, 16, "اختیاری"),
    ("effective_from", "تاریخ اعمال قیمت", False, 14, "شمسی مثل ۱۴۰۵/۰۷/۰۱؛ خالی = امروز"),
    ("parent_sku", "کد کالای اصلی", False, 16, "برای تنوع (مثلاً سری وارداتی دیگر)"),
    ("variant_label", "عنوان تنوع", False, 18, "مثلاً «نسخه امارات»"),
    ("origin", "مبدأ", False, 10, "داخلی یا وارداتی"),
    ("origin_country", "کشور مبدأ", False, 14, "برای کالای وارداتی"),
    ("export_eligible", "قابل صادرات", False, 11, "بله / خیر"),
    ("hs_code", "کد تعرفه گمرکی", False, 14, ""),
    ("description", "توضیحات", False, 30, ""),
    ("specs", "مشخصات", False, 34, "ویژگی: مقدار؛ ویژگی: مقدار — مثلاً «رنگ: مشکی؛ گارانتی: ۱۸ ماه»"),
    ("warehouse", "انبار", False, 18, "نام انبار از فهرست"),
    ("quantity", "موجودی", False, 10, "طبق «نحوه موجودی» هنگام بارگذاری"),
    ("reorder_point", "نقطه سفارش", False, 11, "هشدار کمبود موجودی"),
]
KEYS = [c[0] for c in COLUMNS]
SAMPLE = [
    ["PS5-AE", "کنسول PS5", "کنسول بازی", "Sony", "دستگاه", "", 4.5, 480000000, 0, 420000000, "1405/07/01", "", "", "وارداتی", "امارات",
     "خیر", "95045000", "نسخه استاندارد", "رنگ: سفید؛ گارانتی: ۱۸ ماه", "", 12, 3],
    ["PS5-CN", "کنسول PS5", "کنسول بازی", "Sony", "دستگاه", "", 4.5, 455000000, 0, 401000000, "1405/10/01", "PS5-AE", "سری واردات چین",
     "وارداتی", "چین", "خیر", "95045000", "", "رنگ: سفید؛ گارانتی: ۱۲ ماه", "", 8, 2],
    ["OIL-4L", "روغن مایع ۴ لیتری", "مواد غذایی", "", "عدد", "6260000000001", 3.8, 6200000, 20, "", "", "", "", "وارداتی", "عراق",
     "خیر", "15079010", "یارانه‌ای", "", "", 300, 50],
]


# ---------------------------------------------------------------------------
# قالب و خروجی
# ---------------------------------------------------------------------------
def _style_sheet(ws, warehouses: list[str]) -> None:
    ws.sheet_view.rightToLeft = True
    head = PatternFill("solid", fgColor="0D9488")
    for i, (_key, title, required, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=title + (" *" if required else ""))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"
    letter = {key: ws.cell(row=1, column=i).column_letter for i, key in enumerate(KEYS, start=1)}
    validations = [
        ("unit", "=مقادیر!$A$2:$A$%d" % (len(UNITS) + 1)),
        ("origin", '"داخلی,وارداتی"'),
        ("export_eligible", '"بله,خیر"'),
    ]
    if warehouses:
        validations.append(("warehouse", "=مقادیر!$C$2:$C$%d" % (len(warehouses) + 1)))
    for key, formula in validations:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(f"{letter[key]}2:{letter[key]}{MAX_ROWS + 1}")
        ws.add_data_validation(dv)


def _values_sheet(wb: Workbook, warehouses: list[str]) -> None:
    ws = wb.create_sheet("مقادیر")
    ws.sheet_view.rightToLeft = True
    for col, (title, values) in enumerate((("واحد", UNITS), ("مبدأ", ["داخلی", "وارداتی"]), ("انبار", warehouses)), start=1):
        ws.cell(row=1, column=col, value=title).font = Font(bold=True)
        for r, v in enumerate(values, start=2):
            ws.cell(row=r, column=col, value=v)
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 22


def _guide_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("راهنما")
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    lines = [
        ("راهنمای قالب کاتالوگ پسینه", ""),
        ("", ""),
        ("هر ردیف", "یک کالا در یک انبار. برای ثبت موجودی یک کالا در چند انبار، ردیف را با همان «کد کالا» تکرار کنید."),
        ("کالای موجود", "اگر «کد کالا» از قبل در سامانه باشد، مشخصاتش به‌روز می‌شود؛ وگرنه کالای جدید ساخته می‌شود."),
        ("قیمت", "با «تاریخ اعمال قیمت»، قیمت از آن روز به بعد معتبر می‌شود؛ سفارش‌های قبلی قیمت خودشان را نگه می‌دارند."),
        ("یارانه", "قیمت نهایی فروش = قیمت پایه × (۱ − درصد یارانه / ۱۰۰). برای کالای بدون یارانه خالی یا ۰ بگذارید."),
        ("تنوع", "برای سری دیگری از همان کالا (مثلاً مبدأ یا مشخصات متفاوت) کد جدید بدهید و «کد کالای اصلی» را پر کنید."),
        ("موجودی", "هنگام بارگذاری انتخاب می‌کنید: نادیده گرفتن، تنظیم موجودی به عدد فایل، یا افزودن به موجودی (ورود کالا)."),
        ("ارقام", "ارقام فارسی و جداکننده هزارگان (۱٬۲۵۰٬۰۰۰ یا 1,250,000) قابل قبول است."),
        ("", ""),
    ]
    for r, (a, b) in enumerate(lines, start=1):
        ws.cell(row=r, column=1, value=a).font = Font(bold=True, size=14 if r == 1 else 11)
        ws.cell(row=r, column=2, value=b).alignment = Alignment(wrap_text=True)
    r = len(lines) + 1
    ws.cell(row=r, column=1, value="ستون").font = Font(bold=True)
    ws.cell(row=r, column=2, value="توضیح").font = Font(bold=True)
    for _key, title, required, _, hint in COLUMNS:
        r += 1
        ws.cell(row=r, column=1, value=title + (" (اجباری)" if required else ""))
        ws.cell(row=r, column=2, value=hint)


def _workbook(warehouses: list[str]) -> tuple[Workbook, object]:
    wb = Workbook()
    ws = wb.active
    ws.title = "کالاها"
    _style_sheet(ws, warehouses)
    _guide_sheet(wb)
    _values_sheet(wb, warehouses)
    return wb, ws


def _bytes(wb: Workbook) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def template(db: Session) -> bytes:
    warehouses = [w.name for w in db.query(Warehouse).order_by(Warehouse.id).all()]
    wb, ws = _workbook(warehouses)
    for r, row in enumerate(SAMPLE, start=2):
        row = list(row)
        row[KEYS.index("warehouse")] = warehouses[0] if warehouses else ""
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    return _bytes(wb)


def _jalali(d: date | None) -> str:
    if not d:
        return ""
    jy, jm, jd = gregorian_to_jalali(d.year, d.month, d.day)
    return f"{jy}/{jm:02d}/{jd:02d}"


def export(db: Session, user: User) -> bytes:
    warehouses = {w.id: w.name for w in db.query(Warehouse).order_by(Warehouse.id).all()}
    wb, ws = _workbook(list(warehouses.values()))
    today = catalog_service.today_local()
    products = db.query(Product).order_by(Product.parent_id.isnot(None), Product.id).all()
    skus = {p.id: p.sku for p in products}
    current = catalog_service.prices_at(db, [p.id for p in products], today)
    manager_wh = user.warehouse_id if getattr(user.role, "value", user.role) == "warehouse_manager" else None
    r = 2
    for p in products:
        price = current.get(p.id)
        invs = [i for i in p.inventories if manager_wh is None or i.warehouse_id == manager_wh] or [None]
        if manager_wh is not None and invs == [None]:
            continue
        for inv in invs:
            values = {
                "sku": p.sku, "name": p.name, "category": p.category, "brand": p.brand, "unit": p.unit, "barcode": p.barcode,
                "weight_kg": p.weight_kg, "price": price.base_price if price else p.price,
                "subsidy_percent": price.subsidy_percent if price else 0, "cost_price": price.cost_price if price else None,
                "effective_from": _jalali(price.effective_from) if price else "", "parent_sku": skus.get(p.parent_id, ""),
                "variant_label": p.variant_label, "origin": ORIGIN_LABELS.get(p.origin or "domestic"), "origin_country": p.origin_country,
                "export_eligible": "بله" if p.export_eligible else "خیر", "hs_code": p.hs_code, "description": p.description,
                "specs": "؛ ".join(f"{k}: {v}" for k, v in (p.specs or [])),
                "warehouse": warehouses.get(inv.warehouse_id, "") if inv else "", "quantity": inv.quantity if inv else None,
                "reorder_point": inv.reorder_point if inv else None,
            }
            for c, key in enumerate(KEYS, start=1):
                ws.cell(row=r, column=c, value=values[key])
            r += 1
    return _bytes(wb)


# ---------------------------------------------------------------------------
# ورود
# ---------------------------------------------------------------------------
_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def _text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = normalize_persian_text(str(v)).strip()
    return s or None


def _number(v, field_name: str, errors: list[str], *, integer=False, low=None, high=None):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        n = float(v)
    else:
        s = str(v).translate(_DIGITS).replace("٬", "").replace(",", "").replace("٫", ".").replace("%", "").strip()
        try:
            n = float(s)
        except ValueError:
            errors.append(f"«{field_name}» عدد معتبر نیست ({v}).")
            return None
    if integer and not float(n).is_integer():
        errors.append(f"«{field_name}» باید عدد صحیح باشد.")
        return None
    if (low is not None and n < low) or (high is not None and n > high):
        errors.append(f"«{field_name}» خارج از محدوده مجاز است.")
        return None
    return int(n) if integer else n


def _date(v, errors: list[str]) -> date | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.match(r"^\s*(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})\s*$", str(v).translate(_DIGITS))
    if not m:
        errors.append(f"تاریخ «{v}» معتبر نیست (مثل ۱۴۰۵/۰۷/۰۱).")
        return None
    y, mo, d = map(int, m.groups())
    try:
        if y < 1700:                              # شمسی
            if not (1 <= mo <= 12 and 1 <= d <= (31 if mo <= 6 else 30)):
                raise ValueError
            return date(*jalali_to_gregorian(y, mo, d))
        return date(y, mo, d)
    except ValueError:
        errors.append(f"تاریخ «{v}» معتبر نیست.")
        return None


def _specs(v) -> list[list[str]] | None:
    text = _text(v)
    if not text:
        return None
    pairs = []
    for part in re.split(r"[؛;\n]", text):
        if ":" in part or "：" in part:
            k, _, val = part.replace("：", ":").partition(":")
            if k.strip():
                pairs.append([k.strip(), val.strip()])
        elif part.strip():
            pairs.append([part.strip(), ""])
    return pairs or None


@dataclass
class Row:
    line: int
    data: dict
    errors: list[str] = field(default_factory=list)
    action: str = ""


def _read(content: bytes) -> list[Row]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ConflictError("فایل اکسل خوانده نشد؛ قالب ‎.xlsx‎ سامانه را دانلود و پر کنید.")
    ws = wb["کالاها"] if "کالاها" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    titles = {re.sub(r"\s*\*$", "", str(h or "")).strip(): i for i, h in enumerate(header)}
    missing = [t for k, t, req, *_ in COLUMNS if req and t not in titles]
    if missing:
        raise ConflictError(f"ستون‌های اجباری در فایل نیست: {'، '.join(missing)}. از قالب سامانه استفاده کنید.")
    index = {k: titles.get(t) for k, t, *_ in COLUMNS}
    result = []
    for n, values in enumerate(rows, start=2):
        if not values or all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        if len(result) >= MAX_ROWS:
            raise ConflictError(f"حداکثر {MAX_ROWS} ردیف در هر بارگذاری.")
        result.append(Row(line=n, data={k: (values[i] if i is not None and i < len(values) else None) for k, i in index.items()}))
    if not result:
        raise ConflictError("فایل هیچ ردیف کالایی ندارد.")
    return result


def import_catalog(db: Session, user: User, content: bytes, *, stock_mode: str = "ignore", dry_run: bool = True) -> dict:
    """stock_mode: ignore | set (موجودی = عدد فایل) | add (افزودن به موجودی، مثل ورود کالا)."""
    role = getattr(user.role, "value", user.role)
    if role not in ("org_admin", "warehouse_manager"):
        raise ForbiddenError("ورود کاتالوگ فقط برای ادمین سازمان و مدیر انبار است.")
    rows = _read(content)
    warehouses = {normalize_persian_text(w.name).strip(): w for w in db.query(Warehouse).all()}
    existing = {p.sku: p for p in db.query(Product).all()}
    today = catalog_service.today_local()
    file_skus: dict[str, Row] = {}

    # ۱) اعتبارسنجی همه ردیف‌ها
    for row in rows:
        d, e = row.data, row.errors
        sku, name = _text(d["sku"]), _text(d["name"])
        if not sku:
            e.append("«کد کالا» خالی است.")
        if not name and sku not in existing and sku not in file_skus:      # ردیف تکراری همان کالا (انبار دیگر) نام لازم ندارد
            e.append("«نام کالا» خالی است.")
        price = _number(d["price"], "قیمت پایه", e, low=0)
        if price is None and sku not in existing and sku not in file_skus:
            e.append("«قیمت پایه» خالی است.")
        subsidy = _number(d["subsidy_percent"], "درصد یارانه", e, low=0, high=99.99) or 0
        origin_text = _text(d["origin"])
        origin = {"داخلی": "domestic", "وارداتی": "imported", "domestic": "domestic", "imported": "imported"}.get(origin_text or "داخلی")
        if origin is None:
            e.append("«مبدأ» باید «داخلی» یا «وارداتی» باشد.")
        export_text = _text(d["export_eligible"])
        if export_text not in (None, "بله", "خیر", "yes", "no", "1", "0"):
            e.append("«قابل صادرات» باید «بله» یا «خیر» باشد.")
        warehouse = None
        wh_name = _text(d["warehouse"])
        if wh_name:
            warehouse = warehouses.get(wh_name)
            if warehouse is None:
                e.append(f"انبار «{wh_name}» در سامانه نیست.")
            elif role == "warehouse_manager" and warehouse.id != user.warehouse_id:
                e.append("فقط موجودی انبار خودتان را می‌توانید وارد کنید.")
        quantity = _number(d["quantity"], "موجودی", e, integer=True, low=0)
        if quantity is not None and stock_mode != "ignore" and warehouse is None and not wh_name:
            if role == "warehouse_manager":
                warehouse = db.get(Warehouse, user.warehouse_id)
            else:
                e.append("برای ثبت موجودی، «انبار» را مشخص کنید.")
        unit = _text(d["unit"])
        row.data = {
            "sku": sku, "name": name, "category": _text(d["category"]), "brand": _text(d["brand"]), "unit": unit,
            "barcode": _text(d["barcode"]), "weight_kg": _number(d["weight_kg"], "وزن", e, low=0.001),
            "price": price, "subsidy_percent": subsidy, "cost_price": _number(d["cost_price"], "بهای تمام‌شده", e, low=0),
            "effective_from": _date(d["effective_from"], e) or today, "parent_sku": _text(d["parent_sku"]),
            "variant_label": _text(d["variant_label"]), "origin": origin or "domestic", "origin_country": _text(d["origin_country"]),
            "export_eligible": export_text in ("بله", "yes", "1"), "hs_code": _text(d["hs_code"]),
            "description": _text(d["description"]), "specs": _specs(d["specs"]), "warehouse": warehouse,
            "quantity": quantity, "reorder_point": _number(d["reorder_point"], "نقطه سفارش", e, integer=True, low=0),
        }
        if sku:
            first = file_skus.setdefault(sku, row)
            row.action = "stock" if first is not row else ("update" if sku in existing else "create")
    # کالای اصلی باید در سامانه یا همین فایل باشد (و خودش تنوع نباشد)
    for row in rows:
        parent = row.data.get("parent_sku")
        if parent and row.action in ("create", "update"):
            if parent == row.data["sku"]:
                row.errors.append("کالا نمی‌تواند تنوع خودش باشد.")
            elif parent not in existing and parent not in file_skus:
                row.errors.append(f"کالای اصلی «{parent}» نه در سامانه هست و نه در این فایل.")
            elif (parent in file_skus and file_skus[parent].data.get("parent_sku")) or (parent in existing and existing[parent].parent_id):
                row.errors.append(f"«{parent}» خودش تنوع است؛ کالای اصلی را انتخاب کنید.")

    errors = [{"line": r.line, "sku": r.data.get("sku"), "messages": r.errors} for r in rows if r.errors]
    report = {
        "rows": len(rows), "create": sum(1 for r in rows if r.action == "create" and not r.errors),
        "update": sum(1 for r in rows if r.action == "update" and not r.errors),
        "stock_rows": sum(1 for r in rows if not r.errors and stock_mode != "ignore" and r.data.get("quantity") is not None),
        "errors": errors[:200], "error_count": len(errors), "applied": False, "stock_mode": stock_mode,
        "preview": [{"line": r.line, "sku": r.data.get("sku"), "name": r.data.get("name"), "action": r.action,
                     "price": catalog_service.final_price(r.data["price"], r.data["subsidy_percent"]) if r.data.get("price") is not None else None,
                     "warehouse": r.data["warehouse"].name if r.data.get("warehouse") else None, "quantity": r.data.get("quantity"),
                     "ok": not r.errors} for r in rows[:50]],
    }
    if dry_run or errors:
        return report

    # ۲) ثبت (همه یا هیچ)
    try:
        by_sku: dict[str, Product] = dict(existing)
        ordered = sorted((r for r in rows if r.action in ("create", "update")), key=lambda r: bool(r.data["parent_sku"]))
        for row in ordered:
            d = row.data
            product = by_sku.get(d["sku"])
            fields = {k: d[k] for k in ("category", "brand", "barcode", "description", "origin", "origin_country",
                                        "export_eligible", "hs_code", "specs", "variant_label") if d[k] is not None or k in ("export_eligible",)}
            if d["name"]:
                fields["name"] = d["name"]
            if d["unit"]:
                fields["unit"] = d["unit"]
            if d["weight_kg"]:
                fields["weight_kg"] = d["weight_kg"]
            if d["parent_sku"]:
                fields["parent_id"] = by_sku[d["parent_sku"]].id
            if product is None:
                product = Product(organization_id=user.organization_id, sku=d["sku"], name=d["name"], price=0,
                                  weight_kg=d["weight_kg"] or 1.0)
                db.add(product)
                db.flush()
                by_sku[d["sku"]] = product
            catalog_service.apply_fields(product, fields)
            if d["price"] is not None:
                catalog_service.set_price(db, user, product, effective_from=d["effective_from"], base_price=d["price"],
                                          subsidy_percent=d["subsidy_percent"], cost_price=d["cost_price"], note="ورود از اکسل")
        for row in rows:
            d = row.data
            if d.get("warehouse") is None:
                continue
            product = by_sku[d["sku"]]
            inv = db.query(Inventory).filter(Inventory.product_id == product.id, Inventory.warehouse_id == d["warehouse"].id) \
                .with_for_update().first()
            if inv is None:
                inv = Inventory(organization_id=product.organization_id, product_id=product.id, warehouse_id=d["warehouse"].id,
                                quantity=0, reorder_point=d["reorder_point"] if d["reorder_point"] is not None else 10)
                db.add(inv)
                db.flush()
            elif d["reorder_point"] is not None:
                inv.reorder_point = d["reorder_point"]
            if stock_mode != "ignore" and d["quantity"] is not None:
                if stock_mode == "set":
                    stock_ledger.set_balance(db, inv, d["quantity"], "adjustment", user=user, reason="تنظیم موجودی از فایل اکسل")
                elif d["quantity"] > 0:
                    stock_ledger.record(db, inv, d["quantity"], "receipt", user=user, reason="ورود کالا از فایل اکسل")
        db.commit()
    except Exception:
        db.rollback()
        raise
    report["applied"] = True
    return report
